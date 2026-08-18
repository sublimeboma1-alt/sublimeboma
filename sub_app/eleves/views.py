import json
from io import BytesIO
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.http.multipartparser import MultiPartParser
from django.utils import timezone
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .models import Classe, Eleve, NiveauClasse, Sexe, StatutEleve
from .serializers import clean_eleve_payload, serialize_annee_scolaire, serialize_classe, serialize_eleve, serialize_reference


def parse_json_body(request):
    try:
        return json.loads(request.body.decode('utf-8') or '{}')
    except json.JSONDecodeError:
        raise ValidationError({'detail': 'Donnees JSON invalides.'})


def parse_request_data(request):
    if request.content_type and request.content_type.startswith('multipart/form-data'):
        if request.method == 'POST':
            data, files = request.POST.copy(), request.FILES
        else:
            data, files = MultiPartParser(request.META, request, request.upload_handlers, request.encoding).parse()
            data = data.copy()
        if files.get('photo_file'):
            data['photo_file'] = files['photo_file']
        return data
    return parse_json_body(request)


def validation_error_response(error):
    detail = getattr(error, 'message_dict', None) or {'detail': error.messages}
    return JsonResponse({'errors': detail}, status=400)


def annee_model():
    from sub_app.frais_scolaires.models import AnneeScolaire
    return AnneeScolaire


def filter_by_annee_scolaire(queryset, annee_id):
    AnneeScolaire = annee_model()
    if annee_id == 'toutes':
        return queryset, None
    if annee_id:
        try:
            annee = AnneeScolaire.objects.get(id=annee_id)
        except AnneeScolaire.DoesNotExist:
            return queryset.none(), None
    else:
        annee = AnneeScolaire.objects.filter(est_active=True).first()
    return (queryset.filter(annee_scolaire=annee), annee) if annee else (queryset, None)


def get_filtered_eleves(params):
    queryset = Eleve.objects.select_related('annee_scolaire', 'classe', 'classe__niveau', 'sexe', 'statut').all()
    search = params.get('search', '').strip()
    classe = params.get('classe_id', params.get('classe', '')).strip()
    niveau = params.get('niveau', '').strip()
    statut = params.get('statut', '').strip()
    sexe = params.get('sexe', '').strip()
    queryset, selected_annee = filter_by_annee_scolaire(queryset, params.get('annee_scolaire', params.get('annee', '')).strip())
    if search:
        queryset = queryset.filter(Q(matricule__icontains=search) | Q(nom__icontains=search) | Q(post_nom__icontains=search) | Q(prenom__icontains=search) | Q(telephone__icontains=search) | Q(email__icontains=search))
    if classe:
        queryset = queryset.filter(classe_id=classe)
    if niveau:
        queryset = queryset.filter(classe__niveau_id=niveau)
    if statut:
        queryset = queryset.filter(statut_id=statut)
    if sexe:
        queryset = queryset.filter(sexe_id=sexe)
    return queryset, selected_annee


def get_identity():
    from sub_app.noyau.models import IdentiteEtablissement
    return IdentiteEtablissement.active()


def auto_fit_columns(worksheet):
    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(len(str(cell.value or '')) for cell in column_cells)
        worksheet.column_dimensions[letter].width = min(width + 3, 34)


def style_header(worksheet, row_number):
    fill = PatternFill(start_color='FF0C5A5B', end_color='FF0C5A5B', fill_type='solid')
    font = Font(color='FFFFFFFF', bold=True)
    border = Border(bottom=Side(style='thin', color='FFD2DDE6'))
    for cell in worksheet[row_number]:
        cell.fill, cell.font, cell.border = fill, font, border
        cell.alignment = Alignment(horizontal='center', vertical='center')


def filter_rows(params, selected_annee):
    niveau = NiveauClasse.objects.filter(code=params.get('niveau', '').strip()).first()
    classe_id = params.get('classe_id', params.get('classe', '')).strip()
    classe = Classe.objects.filter(id=classe_id).first() if classe_id else None
    statut = StatutEleve.objects.filter(code=params.get('statut', '').strip()).first()
    sexe = Sexe.objects.filter(code=params.get('sexe', '').strip()).first()
    return [
        ['Annee scolaire', selected_annee.annee if selected_annee else 'Toutes les annees'],
        ['Recherche', params.get('search', '').strip() or 'Toutes'],
        ['Niveau', str(niveau) if niveau else 'Tous'],
        ['Classe', str(classe) if classe else 'Toutes'],
        ['Statut', str(statut) if statut else 'Tous'],
        ['Sexe', str(sexe) if sexe else 'Tous'],
    ]


def report_header(ws, title, identity, selected_annee, params, last_col):
    for row in [1, 2, 3, 4]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
    ws.cell(1, 1, identity.nom if identity else 'Etablissement')
    ws.cell(2, 1, identity.espace if identity else '')
    ws.cell(3, 1, ' | '.join(v for v in [getattr(identity, 'adresse', ''), getattr(identity, 'telephone', ''), getattr(identity, 'email', '')] if v))
    ws.cell(4, 1, title)
    ws.cell(1, 1).font = Font(size=16, bold=True, color='FF14213D')
    ws.cell(2, 1).font = Font(size=12, bold=True, color='FF0C5A5B')
    ws.cell(3, 1).font = Font(size=10, color='FF536271')
    ws.cell(4, 1).font = Font(size=13, bold=True, color='FF14213D')
    ws.append([])
    ws.append(['Date export', timezone.localtime().strftime('%Y-%m-%d %H:%M')])
    ws.append(['Filtres pris en compte', ''])
    for row in filter_rows(params, selected_annee):
        ws.append(row)
    fill = PatternFill(start_color='FFEFF6F6', end_color='FFEFF6F6', fill_type='solid')
    for row in range(6, 14):
        ws.cell(row, 1).font = Font(bold=True, color='FF14213D')
        ws.cell(row, 1).fill = fill
    ws.append([])
    return 15


@login_required
@require_http_methods(['GET'])
def classes_list(request):
    classes = Classe.objects.select_related('niveau', 'classe_maternel', 'classe_primaire', 'classe_humanite', 'section').all()
    niveau = request.GET.get('niveau', '').strip()
    if niveau:
        classes = classes.filter(niveau_id=niveau)
    return JsonResponse({'results': [serialize_classe(classe) for classe in classes]})


@login_required
@require_http_methods(['GET'])
def references_view(request):
    annees = annee_model().objects.all()
    active = next((annee for annee in annees if annee.est_active), None)
    return JsonResponse({'annees_scolaires': [serialize_annee_scolaire(a) for a in annees], 'annee_active_id': active.id if active else None, 'niveaux': [serialize_reference(n) for n in NiveauClasse.objects.filter(est_actif=True)], 'sexes': [serialize_reference(s) for s in Sexe.objects.filter(est_actif=True)], 'statuts': [serialize_reference(s) for s in StatutEleve.objects.filter(est_actif=True)]})


@login_required
@csrf_protect
@require_http_methods(['GET', 'POST'])
def eleves_list_create(request):
    if request.method == 'POST':
        try:
            data = clean_eleve_payload(parse_request_data(request))
            eleve = Eleve(**data, created_by=request.user)
            # clean_fields() est plus leger que full_clean() (pas de validate_unique)
            # Le matricule est genere dans save() si absent.
            eleve.clean_fields()
            eleve.save()
        except ValidationError as error:
            return validation_error_response(error)
        return JsonResponse(serialize_eleve(eleve), status=201)
    queryset, active_annee = get_filtered_eleves(request.GET)
    return JsonResponse({'annee_scolaire': serialize_annee_scolaire(active_annee) if active_annee else None, 'results': [serialize_eleve(e) for e in queryset]})


@login_required
@require_http_methods(['GET'])
def export_eleves_xlsx(request):
    queryset, active_annee = get_filtered_eleves(request.GET)
    eleves = list(queryset)
    identity = get_identity()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Eleves'
    headers = ['Matricule', 'Nom', 'Post-nom', 'Prenom', 'Nom complet', 'Sexe', 'Date naissance', 'Lieu naissance', 'Classe', 'Niveau', 'Annee scolaire', 'Statut', 'MASP', 'Telephone', 'Email', 'Adresse', "Date d'inscription", 'Photo']
    header_row = report_header(ws, 'Rapport des eleves', identity, active_annee, request.GET, len(headers))
    ws.append(headers)
    for e in eleves:
        ws.append([e.matricule, e.nom, e.post_nom, e.prenom, f'{e.nom} {e.post_nom} {e.prenom}'.strip(), e.get_sexe_display(), e.date_naissance.isoformat() if e.date_naissance else '', e.lieu_de_naissance or '', str(e.classe) if e.classe else '', str(e.classe.niveau) if e.classe else '', str(e.annee_scolaire) if e.annee_scolaire else '', e.get_statut_display(), 'Oui' if e.est_masp else 'Non', e.telephone or '', e.email or '', e.adresse or '', e.date_inscription.isoformat() if e.date_inscription else '', e.photo.url if e.photo else ''])
    style_header(ws, header_row)
    ws.freeze_panes = f'A{header_row + 1}'
    auto_fit_columns(ws)

    stats = wb.create_sheet('Statistiques')
    stats_header = report_header(stats, 'Statistiques des eleves', identity, active_annee, request.GET, 4)
    stats.append(['Indicateur', 'Valeur'])
    rows = [('Total eleves', len(eleves)), ('Actifs', sum(1 for e in eleves if e.statut_id == 'actif')), ('Inactifs', sum(1 for e in eleves if e.statut_id == 'inactif')), ('Suspendus', sum(1 for e in eleves if e.statut_id == 'suspendu')), ('Diplomes', sum(1 for e in eleves if e.statut_id == 'diplome')), ('Garcons', sum(1 for e in eleves if e.sexe_id == 'M')), ('Filles', sum(1 for e in eleves if e.sexe_id == 'F')), ('Membres MASP', sum(1 for e in eleves if e.est_masp))]
    for row in rows:
        stats.append(row)
    stats.append([])
    stats.append(['Classe', 'Effectif'])
    counts = {}
    for e in eleves:
        key = str(e.classe) if e.classe else 'Non assigne'
        counts[key] = counts.get(key, 0) + 1
    for key, count in sorted(counts.items()):
        stats.append([key, count])
    style_header(stats, stats_header)
    auto_fit_columns(stats)

    output = BytesIO()
    wb.save(output)
    filename = f'eleves_{timezone.localtime().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
@csrf_protect
@require_http_methods(['GET', 'PATCH', 'PUT', 'DELETE'])
def eleve_detail(request, eleve_id):
    try:
        eleve = Eleve.objects.select_related('annee_scolaire', 'classe', 'classe__niveau', 'sexe', 'statut').get(id=eleve_id)
    except Eleve.DoesNotExist:
        return JsonResponse({'detail': 'Eleve introuvable.'}, status=404)
    if request.method == 'GET':
        return JsonResponse(serialize_eleve(eleve))
    if request.method == 'DELETE':
        eleve.delete()
        return JsonResponse({}, status=204)
    try:
        data = clean_eleve_payload(parse_request_data(request), partial=True)
        for field, value in data.items():
            setattr(eleve, field, value)
        # clean_fields() est plus leger que full_clean() (pas de validate_unique)
        eleve.clean_fields()
        eleve.save()
    except ValidationError as error:
        return validation_error_response(error)
    return JsonResponse(serialize_eleve(eleve))
